"""
Document processing module for Social Support AI Workflow

Handles multimodal document processing including:
- PDF text extraction
- Image text recognition (OCR)
- Excel/CSV data parsing
- Bank statement processing
- Emirates ID information extraction
"""
import os
import json
from typing import Dict, List, Optional, Union, Any
from pathlib import Path
import pandas as pd
from PIL import Image
import pytesseract
import numpy as np
from PyPDF2 import PdfReader
import openpyxl
import sys

# Add OpenCV import for image preprocessing
try:
    import cv2
except ImportError:
    cv2 = None

# Add project root to path
sys.path.append(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))

# Import logging configuration
from src.utils.logging_config import get_logger

# Setup logging
logger = get_logger("document_processor")

from config import settings


class DocumentProcessor:
    """Handles processing of various document types for data extraction"""
    
    def __init__(self):
        self.supported_formats = {
            'pdf': self._process_pdf,
            'jpg': self._process_image,
            'jpeg': self._process_image,
            'png': self._process_image,
            'xlsx': self._process_excel,
            'xls': self._process_excel,
            'docx': self._process_docx,
            'doc': self._process_docx,
            'csv': self._process_csv
        }
    
    def process_document(self, file_path: str, document_type: str) -> Dict[str, Any]:
        """
        Process a document and extract relevant information
        
        Args:
            file_path: Path to the document file
            document_type: Type of document (bank_statement, emirates_id, resume, etc.)
            
        Returns:
            Dictionary containing extracted information
        """
        try:
            logger.info(f"Starting document processing: {file_path} (type: {document_type})")
            
            # Check if file exists
            if not os.path.exists(file_path):
                logger.error(f"File not found: {file_path}")
                raise FileNotFoundError(f"File not found: {file_path}")
            
            # Get file info
            file_size = os.path.getsize(file_path)
            logger.info(f"File size: {file_size} bytes")
            
            file_extension = Path(file_path).suffix.lower().lstrip('.')
            logger.info(f"File extension: {file_extension}")
            
            if file_extension not in self.supported_formats:
                logger.error(f"Unsupported file format: {file_extension}")
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            logger.info(f"Processing {document_type} document: {file_path}")
            
            # Extract raw data using appropriate processor
            logger.info(f"Using processor for extension: {file_extension}")
            raw_data = self.supported_formats[file_extension](file_path)
            logger.info(f"Raw data extraction completed. Data type: {raw_data.get('type', 'unknown')}")
            
            # Apply document-specific processing
            logger.info(f"Applying document-specific processing for: {document_type}")
            processed_data = self._apply_document_specific_processing(
                raw_data, document_type
            )
            logger.info(f"Document-specific processing completed")
            
            result = {
                'status': 'success',
                'document_type': document_type,
                'raw_data': raw_data,
                'processed_data': processed_data,
                'metadata': {
                    'file_path': file_path,
                    'file_size': file_size,
                    'file_extension': file_extension
                }
            }
            
            logger.info(f"Document processing completed successfully for: {file_path}")
            return result
            
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {str(e)}", exc_info=True)
            return {
                'status': 'error',
                'error': str(e),
                'document_type': document_type,
                'file_path': file_path
            }
    
    def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from PDF documents"""
        try:
            reader = PdfReader(file_path)
            text_content = []
            
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                text_content.append({
                    'page': page_num + 1,
                    'text': text
                })
            
            return {
                'type': 'pdf',
                'pages': len(reader.pages),
                'content': text_content,
                'full_text': '\n'.join([page['text'] for page in text_content])
            }
            
        except Exception as e:
            logger.error(f"Error processing PDF {file_path}: {e}")
            raise
    
    def _process_image(self, file_path: str) -> Dict[str, Any]:
        """Extract text from images using OCR"""
        try:
            logger.info(f"Starting image processing for: {file_path}")
            
            # Check if OpenCV is available
            if cv2 is None:
                logger.warning("OpenCV (cv2) not available")
                return {
                    'type': 'image',
                    'text': '',
                    'error': 'OpenCV not available - required for image preprocessing',
                    'ocr_confidence': 0,
                    'image_dimensions': None,
                    'detected_text_blocks': 0,
                    'installation_note': 'To enable image processing, install OpenCV: pip install opencv-python'
                }
            
            # Check if tesseract is available
            try:
                import pytesseract
                # Test if tesseract is installed
                version = pytesseract.get_tesseract_version()
                logger.info(f"Tesseract version: {version}")
            except (ImportError, pytesseract.TesseractNotFoundError) as e:
                logger.warning(f"Tesseract not available: {str(e)}")
                return {
                    'type': 'image',
                    'text': '',
                    'error': 'OCR not available - Tesseract not installed',
                    'ocr_confidence': 0,
                    'image_dimensions': None,
                    'detected_text_blocks': 0,
                    'installation_note': 'To enable OCR processing, install Tesseract: brew install tesseract (macOS) or apt-get install tesseract-ocr (Ubuntu)'
                }
            
            # Load and preprocess image
            logger.info(f"Loading image: {file_path}")
            image = cv2.imread(file_path)
            if image is None:
                logger.error(f"Could not load image from {file_path}")
                raise ValueError(f"Could not load image from {file_path}")
            
            logger.info(f"Image loaded successfully. Shape: {image.shape}")
            
            gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
            logger.info("Converted image to grayscale")
            
            # Apply image enhancement for better OCR
            gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
            logger.info("Applied image thresholding")
            
            # Extract text using OCR
            logger.info("Starting OCR text extraction")
            text = pytesseract.image_to_string(gray)
            logger.info(f"OCR text extraction completed. Text length: {len(text)}")
            
            # Get additional OCR data
            logger.info("Getting OCR confidence data")
            ocr_data = pytesseract.image_to_data(gray, output_type=pytesseract.Output.DICT)
            
            confidence_scores = [int(conf) for conf in ocr_data['conf'] if int(conf) > 0]
            avg_confidence = np.mean(confidence_scores) if confidence_scores else 0
            text_blocks = len([t for t in ocr_data['text'] if t.strip()])
            
            logger.info(f"OCR completed. Confidence: {avg_confidence:.2f}, Text blocks: {text_blocks}")
            
            result = {
                'type': 'image',
                'text': text,
                'ocr_confidence': avg_confidence,
                'image_dimensions': image.shape[:2],
                'detected_text_blocks': text_blocks
            }
            
            logger.info(f"Image processing completed successfully for: {file_path}")
            return result
            
        except Exception as e:
            logger.error(f"Error processing image {file_path}: {str(e)}", exc_info=True)
            
            # Return a more helpful error message
            error_msg = str(e)
            if "tesseract" in error_msg.lower():
                error_msg = "OCR service not available. To process images, please install Tesseract OCR."
                installation_help = {
                    'macOS': 'brew install tesseract',
                    'Ubuntu/Debian': 'sudo apt-get install tesseract-ocr',
                    'Windows': 'Download from: https://github.com/UB-Mannheim/tesseract/wiki'
                }
            elif "cv2" in error_msg.lower() or "opencv" in error_msg.lower():
                error_msg = "OpenCV not available. To process images, please install OpenCV."
                installation_help = {
                    'pip': 'pip install opencv-python'
                }
            else:
                installation_help = {}
            
            return {
                'type': 'image',
                'text': '',
                'error': error_msg,
                'ocr_confidence': 0,
                'image_dimensions': None,
                'detected_text_blocks': 0,
                'installation_help': installation_help
            }
    
    def _process_excel(self, file_path: str) -> Dict[str, Any]:
        """Process Excel files and extract tabular data"""
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            sheets_data = {}
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                sheets_data[sheet_name] = {
                    'shape': df.shape,
                    'columns': df.columns.tolist(),
                    'data': df.to_dict('records'),
                    'summary': {
                        'total_rows': len(df),
                        'total_columns': len(df.columns),
                        'numeric_columns': len(df.select_dtypes(include=[np.number]).columns),
                        'text_columns': len(df.select_dtypes(include=['object']).columns)
                    }
                }
            
            return {
                'type': 'excel',
                'sheets': list(excel_file.sheet_names),
                'data': sheets_data
            }
            
        except Exception as e:
            logger.error(f"Error processing Excel file {file_path}: {e}")
            raise
    
    def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """Process Word documents"""
        try:
            doc = Document(file_path)
            
            # Extract text from paragraphs
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract tables if any
            tables_data = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables_data.append(table_data)
            
            return {
                'type': 'docx',
                'paragraphs': paragraphs,
                'full_text': '\n'.join(paragraphs),
                'tables': tables_data,
                'paragraph_count': len(paragraphs),
                'table_count': len(tables_data)
            }
            
        except Exception as e:
            logger.error(f"Error processing DOCX file {file_path}: {e}")
            raise
    
    def _process_csv(self, file_path: str) -> Dict[str, Any]:
        """Process CSV files"""
        try:
            df = pd.read_csv(file_path)
            
            return {
                'type': 'csv',
                'shape': df.shape,
                'columns': df.columns.tolist(),
                'data': df.to_dict('records'),
                'summary': {
                    'total_rows': len(df),
                    'total_columns': len(df.columns),
                    'numeric_columns': len(df.select_dtypes(include=[np.number]).columns),
                    'text_columns': len(df.select_dtypes(include=['object']).columns)
                }
            }
            
        except Exception as e:
            logger.error(f"Error processing CSV file {file_path}: {e}")
            raise
    
    def _apply_document_specific_processing(self, raw_data: Dict[str, Any], document_type: str) -> Dict[str, Any]:
        """Apply document type specific processing logic"""
        
        if document_type == "bank_statement":
            return self._process_bank_statement(raw_data)
        elif document_type == "emirates_id":
            return self._process_emirates_id(raw_data)
        elif document_type == "resume":
            return self._process_resume(raw_data)
        elif document_type == "credit_report":
            return self._process_credit_report(raw_data)
        elif document_type == "assets":
            return self._process_assets_document(raw_data)
        else:
            return {"processed": False, "reason": f"No specific processor for {document_type}"}
    
    def _process_bank_statement(self, raw_data: Dict[str, Any]) -> Dict[str, Any]:
        """Extract financial information from bank statements"""
        # This is a simplified version - in production, you'd use more sophisticated parsing
        text = raw_data.get('full_text', '') or raw_data.get('text', '')
        
        # Extract key financial indicators (simplified pattern matching)
        import re
        
        # Find amounts (simplified regex)
        amounts = re.findall(r'\b\d{1,3}(?:,\d{3})*(?:\.\d{2})?\b', text)
        
        return {
            'financial_summary': {
                'detected_amounts': amounts[:10],  # First 10 amounts found
                'potential_income_indicators': self._find_income_patterns(text),
                'potential_expense_indicators': self._find_expense_patterns(text)
            },
            'text_length': len(text),
            'contains_financial_data': len(amounts) > 0
        }
    
    def _process_emirates_id(self, raw_data: Dict[str, Any]) -> Dict[str, Any]:
        """Extract information from Emirates ID"""
        text = raw_data.get('full_text', '') or raw_data.get('text', '')
        
        # Extract Emirates ID specific information (simplified)
        import re
        
        # Emirates ID number pattern (simplified)
        emirates_id_pattern = r'\b\d{3}-\d{4}-\d{7}-\d{1}\b'
        emirates_ids = re.findall(emirates_id_pattern, text)
        
        return {
            'identity_info': {
                'emirates_id_detected': len(emirates_ids) > 0,
                'emirates_id_numbers': emirates_ids,
                'text_confidence': raw_data.get('ocr_confidence', 0)
            }
        }
    
    def _process_resume(self, raw_data: Dict[str, Any]) -> Dict[str, Any]:
        """Extract career information from resume"""
        text = raw_data.get('full_text', '') or raw_data.get('text', '')
        
        # Simple keyword-based analysis
        experience_keywords = ['experience', 'work', 'employment', 'career', 'position']
        education_keywords = ['education', 'degree', 'university', 'college', 'certification']
        
        return {
            'career_info': {
                'has_experience_section': any(keyword in text.lower() for keyword in experience_keywords),
                'has_education_section': any(keyword in text.lower() for keyword in education_keywords),
                'text_length': len(text)
            }
        }
    
    def _process_credit_report(self, raw_data: Dict[str, Any]) -> Dict[str, Any]:
        """Extract credit information"""
        text = raw_data.get('full_text', '') or raw_data.get('text', '')
        
        # Look for credit-related keywords
        credit_keywords = ['credit score', 'payment history', 'debt', 'loan', 'credit card']
        
        return {
            'credit_info': {
                'contains_credit_data': any(keyword in text.lower() for keyword in credit_keywords),
                'text_length': len(text)
            }
        }
    
    def _process_assets_document(self, raw_data: Dict[str, Any]) -> Dict[str, Any]:
        """Process assets and liabilities document"""
        if raw_data['type'] in ['excel', 'csv']:
            # For tabular data, look for asset/liability columns
            data = raw_data.get('data', {})
            
            asset_liability_analysis = {}
            
            if raw_data['type'] == 'excel':
                for sheet_name, sheet_data in data.items():
                    columns = sheet_data.get('columns', [])
                    asset_liability_analysis[sheet_name] = {
                        'has_asset_columns': any('asset' in col.lower() for col in columns),
                        'has_liability_columns': any('liability' in col.lower() or 'debt' in col.lower() for col in columns),
                        'numeric_columns': sheet_data.get('summary', {}).get('numeric_columns', 0)
                    }
            else:
                columns = data.get('columns', [])
                asset_liability_analysis['main'] = {
                    'has_asset_columns': any('asset' in col.lower() for col in columns),
                    'has_liability_columns': any('liability' in col.lower() or 'debt' in col.lower() for col in columns),
                    'numeric_columns': data.get('summary', {}).get('numeric_columns', 0)
                }
            
            return {
                'assets_liabilities_info': asset_liability_analysis
            }
        
        return {'processed': False, 'reason': 'Assets document should be in tabular format'}
    
    def _find_income_patterns(self, text: str) -> List[str]:
        """Find potential income-related patterns in text"""
        income_patterns = ['salary', 'income', 'wage', 'deposit', 'credit']
        found_patterns = []
        
        for pattern in income_patterns:
            if pattern in text.lower():
                found_patterns.append(pattern)
        
        return found_patterns
    
    def _find_expense_patterns(self, text: str) -> List[str]:
        """Find potential expense-related patterns in text"""
        expense_patterns = ['payment', 'withdrawal', 'debit', 'charge', 'fee']
        found_patterns = []
        
        for pattern in expense_patterns:
            if pattern in text.lower():
                found_patterns.append(pattern)
        
        return found_patterns 